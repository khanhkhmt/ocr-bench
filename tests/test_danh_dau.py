from docx import Document

from htr_test.danh_dau import so_khop_tu, them_doan_danh_dau, ty_le_trung


def test_khac_biet_khong_doi_nghia_khong_bi_danh_dau():
    # hamza / alef, ى/ي, dấu nguyên âm, dấu câu dính kèm → coi là trùng
    assert not any(ngo for _, ngo, _ in so_khop_tu("إن الملك يعلم.", "ان المَلك يعلم"))
    assert ty_le_trung("على الامر", "علي الأمر") == 1.0


def test_tu_them_doi_nghia_bi_danh_dau():
    tu = so_khop_tu("أنهم لا يوافقون على قوة دولية", "انهم يوافقون على قوة دولية")
    assert [w for w, ngo, _ in tu if ngo] == ["لا"]
    assert dict((w, alt) for w, ngo, alt in tu if ngo)["لا"] == "(không có)"


def test_docx_to_vang_va_chu_thich(tmp_path):
    doc = Document()
    _, n, m = them_doan_danh_dau(doc, "ووزير الثقافة والاعلام", "ورئيس الثقافة والاعلام", ten_phu="B")
    f = tmp_path / "x.docx"
    doc.save(f)
    d = Document(f)
    assert (n, m) == (3, 1)
    assert sum(1 for r in d.paragraphs[0].runs if r.font.highlight_color is not None) == 1
    assert [c.text for c in d.comments] == ["B đọc: ورئيس"]
