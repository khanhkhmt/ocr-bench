"""Đánh dấu chữ ĐÁNG NGỜ bằng so khớp hai bộ đọc độc lập, và ghi vào DOCX (tô vàng + chú thích cách đọc kia).

Cơ sở (NGHIEN_CUU.md, chủ đề 3 — 240 dòng Omar blind_test, Baseer vs Ketaba): đánh dấu từ mà bộ đọc chính có nhưng bộ
đọc phụ không có → bắt ~78% từ sai, chỉ đánh dấu ~22% số từ; dòng hai bộ gần như trùng → đúng ~98,7%.

    from htr_test.danh_dau import so_khop_tu, them_doan_danh_dau
    them_doan_danh_dau(doc, chinh="…", phu="…", ten_phu="Ketaba")
"""

from __future__ import annotations

import difflib
import re

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_HARAKAT = re.compile(r"[ً-ْٰـ]")
_ALEF = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه"})
_PUNCT = re.compile(r"[\s\.\,\:\;\!\?\(\)\[\]\"'«»،؛؟\-–—]+")


def _key(w: str) -> str:
    """So sánh từ bỏ qua khác biệt KHÔNG đổi nghĩa: dấu nguyên âm, hamza/alef, ى/ي, ة/ه, dấu câu dính kèm."""
    return _PUNCT.sub("", _HARAKAT.sub("", w)).translate(_ALEF)


def so_khop_tu(chinh: str, phu: str) -> list[tuple[str, bool, str]]:
    """→ [(từ của bộ đọc chính, có đáng ngờ không, bộ đọc phụ đọc chỗ đó là gì)].
    Đáng ngờ = từ nằm trong đoạn hai bộ đọc KHÁC nhau (thay / thừa so với bộ phụ)."""
    a, b = chinh.split(), phu.split()
    sm = difflib.SequenceMatcher(None, [_key(w) for w in a], [_key(w) for w in b], autojunk=False)
    out: list[tuple[str, bool, str]] = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        alt = " ".join(b[j1:j2])
        for w in a[i1:i2]:
            out.append((w, tag != "equal", "" if tag == "equal" else (alt or "(không có)")))
    return out


def ty_le_trung(chinh: str, phu: str) -> float:
    """Tỉ lệ từ trùng (0–1) — dòng ≥ 0,97 coi như hai bộ đọc đồng ý."""
    tu = so_khop_tu(chinh, phu)
    return 1.0 if not tu else sum(not d for _, d, _ in tu) / len(tu)


def _rtl(run) -> None:
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))


def them_doan_danh_dau(doc, chinh: str, phu: str, ten_phu: str = "bộ đọc 2", tac_gia: str = "OCR",
                       chu_thich: bool = True):
    """Thêm MỘT đoạn (phải → trái) vào doc: chữ của bộ đọc chính; từ đáng ngờ tô VÀNG, gom từng cụm liền nhau vào một
    chú thích Word "<ten_phu> đọc: …" (chu_thich=False → chỉ tô vàng). Trả về (đoạn, số từ, số từ đáng ngờ)."""
    from ocrbench.docx_export import _set_bidi  # cùng cách đặt đoạn RTL như DOCX A

    p = doc.add_paragraph()
    _set_bidi(p)
    tu = so_khop_tu(chinh, phu)
    i, n_ngo = 0, 0
    while i < len(tu):
        w, ngo, alt = tu[i]
        if not ngo:
            r = p.add_run(("" if i == 0 else " ") + w)
            _rtl(r)
            i += 1
            continue
        runs, j = [], i
        while j < len(tu) and tu[j][1] and tu[j][2] == alt:  # cùng một đoạn lệch → một chú thích
            if j:
                _rtl(p.add_run(" "))
            r = p.add_run(tu[j][0])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            _rtl(r)
            runs.append(r)
            j += 1
        n_ngo += len(runs)
        if chu_thich:
            doc.add_comment(runs, text=f"{ten_phu} đọc: {alt}", author=tac_gia, initials="OCR")
        i = j
    return p, len(tu), n_ngo
